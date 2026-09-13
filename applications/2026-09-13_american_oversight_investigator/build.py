import os,sys,json,subprocess
from pathlib import Path
ROOT=Path(__file__).resolve().parent
(ROOT/'review').mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT.parents[1]))
from docx_header import *
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
REPO=ROOT.parents[1]
standards=['README.md','HEADER_STANDARD.md','docx_header.py','EDUCATION_CONSTANTS.md','CAREER_CONSTANTS.md','RESUME_HEADLINE_SCANABILITY_STANDARD.md','INTELLIGENCE_OSINT_CASE_EVIDENCE.md','standards/document_design_standard.json']
(ROOT/'review'/'provenance.json').write_text(json.dumps({'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=REPO,text=True).strip(),'blobs':{f:subprocess.check_output(['git','rev-parse','HEAD:'+f],cwd=REPO,text=True).strip() for f in standards},'lane':'analyst-intelligence','layout':'Locked helper, repeating full DOCX header; Letter; top 1.55, sides .65, bottom .55; body 10.25; sections 14/6; entries 8/2/4; line 1.05. Separate portal resume uses body contact and no header. Cover 10.5, sides .78, bottom .68.','preflight':'PASS before generation'},indent=2))
def base(branded=True,cover=False):
 d=new_document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
 for st in d.styles:
  if hasattr(st,'font'):
   st.font.name='EB Garamond'
   if st._element.rPr is not None and st._element.rPr.rFonts is not None:
    f=st._element.rPr.rFonts
    for a in list(f.attrib):
     if 'Theme' in a: del f.attrib[a]
 if branded: build_navy_header(d,body_top_margin_inches=1.55,body_bottom_margin_inches=.68 if cover else .55,body_left_margin_inches=.78 if cover else .65,body_right_margin_inches=.78 if cover else .65)
 else:
  s.top_margin=Inches(.65);s.bottom_margin=Inches(.55);s.left_margin=s.right_margin=Inches(.65)
  p=d.add_paragraph();set_run(p.add_run(NAME),size=20,bold=True)
  p=d.add_paragraph();set_run(p.add_run(' | '.join(t for t,u in CONTACT_PARTS)),size=10)
 d.core_properties.author='Troy Hokanson';d.core_properties.last_modified_by='Troy Hokanson'
 return d

def para(d,t,bold=False,before=0,after=4,size=10.25,keep=False):
 p=d.add_paragraph();set_run(p.add_run(t.replace(chr(8217), chr(39))),size=size,bold=bold);set_paragraph_format(p,before=before,after=after,line=1.05);p.paragraph_format.keep_with_next=keep;p.paragraph_format.widow_control=True;return p

def heading(d,t,first=False):
 p=add_section_heading(d,t);set_paragraph_format(p,before=18 if first else 14,after=6,line=1.05);p.paragraph_format.keep_with_next=True
 for r in p.runs:r.font.size=Pt(11)
 return p

def job(d,title,org,dates,bullets=[]):
 p1,p2=add_job_block(d,title,org,dates)
 set_paragraph_format(p1,before=8,after=2,line=1.05);set_paragraph_format(p2,after=4,line=1.05)
 for p in [p1,p2]:p.paragraph_format.keep_with_next=True
 for r in p1.runs:r.font.color.rgb=STEEL;r.font.size=Pt(10.5)
 for t in bullets:
  p=add_bullet(d,t.replace(chr(8217), chr(39)),size=10.25);set_paragraph_format(p,after=2,line=1.05);p.paragraph_format.widow_control=True

def resume(branded):
 d=base(branded)
 para(d,'Investigative Research | Records Analysis | Evidence and Reporting',True,before=18 if branded else 8,size=11)
 para(d,'Investigator with 6.5 years of detective assignments within more than 25 years of policing, including 5.5 years in digital forensics. Experienced in records research, interviews, legal-process drafting, and reports for prosecutors. Remote criminal-justice educator for 18 years.')
 heading(d,'Investigative experience')
 job(d,'Detective / Digital Forensic Examiner','Dakota County Electronic Crimes Task Force, assigned from Lakeville PD','June 2017 - December 2021',[
 'Examined computer and mobile evidence, evaluated findings against case facts, and documented results for investigators and prosecutors. Preserved evidence and maintained chain of custody.',
 'Prepared investigative and forensic reports, coordinated with prosecutors, and explained evidence in courtroom testimony. Distinguished supported findings from investigative leads requiring further work.',
 'Used open-source research to document public social-media profiles, corroborate identities with Accurint, and preserve online evidence for follow-up investigation.',
 'Served as Lakeville’s representative in a task force supporting ten partner agencies and as the department’s digital-forensics subject-matter resource.' ])
 job(d,'Detective / Electronic Crimes Unit','Lakeville Police Department, Minnesota','September 2016 - June 2017',[
 'Drafted search warrants and researched provider records requirements. Built an investigator reference folder containing subpoena, preservation, and search-warrant templates with service-provider guidance.',
 'Conducted digital-evidence examinations and prepared reports to support case development and attorney review.' ])
 job(d,'Police Officer / Investigator (First Investigative Rotation)','Lakeville Police Department, Minnesota','March 2010 - May 2011',[
 'Investigated financial and property crimes through records research, witness and subject interviews, evidence gathering, and reports supporting prosecution.' ])
 heading(d,'Research and technical skills')
 para(d,'LexisNexis Accurint: 6.5 years of direct investigative use across both detective rotations. Investigative research, source comparison, evidence organization, interviews, and written findings. Microsoft Word, Excel, Outlook, and Teams.')
 para(d,'Historical forensic tools: AccessData Forensic Toolkit (FTK), FTK Imager, X-Ways, and Cellebrite. Operational digital-forensics work ended in 2021. Odyssey electronic filing of search warrants.')
 d.add_page_break()
 heading(d,'Additional experience',True)
 job(d,'Independent Professional','Independent, Remote','April 2026 - Present', ['Complete paid AI project work and professional development.'])
 job(d,'Real Estate Consultant','eXp Realty / KW Select, South Metro Minnesota','June 2024 - June 2026')
 job(d,'Adjunct Faculty / Criminal Justice','University of Phoenix, Remote','March 2007 - October 2025',[
 'Taught criminal justice online for 18 years, explaining legal and investigative concepts to students and providing written feedback. Concurrent with law-enforcement service.' ])
 job(d,'Police Officer','Lakeville Police Department, Minnesota','January 2022 - May 2024')
 job(d,'Police Officer / Field Training Officer','Lakeville Police Department, Minnesota','June 2011 - August 2016')
 job(d,'Police Officer / Field Training Officer','Lakeville Police Department, Minnesota','November 1998 - February 2010',[
 'Coauthored a $40,000 Target + Blue grant for automated license plate recognition and coordinated implementation with technical partners.' ])
 heading(d,'Education')
 for degree,school,gpa,year in [
 ('Master of Arts, Police Leadership, Administration and Education','University of St. Thomas, St. Paul, MN','3.94','2005'),
 ('Bachelor of Arts, Criminal Justice, Magna Cum Laude','St. Cloud State University, St. Cloud, MN','3.51','1998'),
 ('Associate of Arts, Criminal Justice, Magna Cum Laude','St. Cloud State University, St. Cloud, MN','3.50','1996')]:
  para(d,degree,True,before=4,after=0,keep=True)
  para(d,school+'\nGPA: '+gpa+'\n'+year,after=2)
 heading(d,'Credentials and training')
 para(d,'Certified Cyber Crime Investigator (CCCI), No. 4793, January 2023. Private-investigator training: 18 hours, 2026 (training only, not licensure).')
 name='Troy_Hokanson_American_Oversight_Resume'+('' if branded else '_ATS')
 d.core_properties.title='Troy Hokanson | American Oversight Investigator | Resume'
 d.save(ROOT/(name+'.docx'))
 (ROOT/(name+'.txt')).write_text('\n'.join(p.text for p in d.paragraphs))
resume(True);resume(False)
d=base(True,True)
para(d,'September 13, 2026',before=18,after=10,size=10.5)
para(d,'Hiring Committee\nAmerican Oversight',after=10,size=10.5)
para(d,'Re: Investigator',True,after=10,size=10.5)
para(d,'Dear Hiring Committee:',after=8,size=10.5)
texts=[
'I am applying for the Investigator position at American Oversight. The work of following a lead through records, testing what those records show, and explaining the findings is familiar to me. My background includes 6.5 years of detective assignments, including 5.5 years of digital-forensics work, within more than 25 years with the Lakeville Police Department.',
'As a detective and digital forensic examiner, I reviewed evidence, researched records, interviewed people, and prepared reports for investigators and prosecutors. That work required careful attention to the limits of the evidence as well as its significance. I also drafted search warrants and built a reference folder of subpoena, preservation, and warrant templates with service-provider information so other investigators could pursue records more consistently.',
'Your investigators turn document review into findings that other people can use. My written products have primarily served criminal investigations and prosecution. I have explained technical evidence in reports and courtroom testimony, and I taught criminal justice remotely for 18 years at the University of Phoenix. Teaching required me to make unfamiliar concepts understandable and to give clear written feedback to people with different levels of experience.',
'Public-records requests and public-facing accountability research would be a new setting for my investigative work. I would need to learn American Oversight’s FOIA practices and research priorities. The opportunity to examine government records carefully and help explain what they establish is what draws me to this position.',
'Thank you for considering my application. I would welcome a conversation about how my experience developing cases and explaining evidence could support your investigations.'
]
for t in texts:para(d,t,after=8,size=10.5)
para(d,'Sincerely,\nTroy Hokanson',before=4,size=10.5)
d.core_properties.title='Troy Hokanson | American Oversight Investigator | Cover Letter'
d.save(ROOT/'Troy_Hokanson_American_Oversight_Cover_Letter.docx')
(ROOT/'Troy_Hokanson_American_Oversight_Cover_Letter.txt').write_text('\n'.join(p.text for p in d.paragraphs))

# Normalize unused Word font defaults without modifying the locked header.
import zipfile,re
for path in ROOT.glob('*.docx'):
 with zipfile.ZipFile(path) as z: contents={n:z.read(n) for n in z.namelist()}
 for n in ['word/styles.xml','word/stylesWithEffects.xml','word/fontTable.xml','word/theme/theme1.xml']:
  if n in contents:
   text=contents[n].decode()
   text=re.sub(r'Calibri Light|Calibri|Arial', 'EB Garamond', text)
   contents[n]=text.encode()
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,data in contents.items(): z.writestr(n,data)
